import pptx
from pptx.util import Pt
from pptx.dml.color import RGBColor
import os
import re
import glob
import shutil
import datetime

output_path = "/Users/refatishrakhemel/Desktop/UNB_MCS/Summer26/RobustML(DirectedStudies)/Model_Extraction_Attacks_Presentation.pptx"
template_path = "/Users/refatishrakhemel/Desktop/UNB_MCS/Summer26/RobustML(DirectedStudies)/Presentation Template 1.pptx"

# -------------------------------------------------------------
# 1. Backup existing file (if it exists)
# -------------------------------------------------------------
if os.path.exists(output_path):
    timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_dir = "/Users/refatishrakhemel/Desktop/UNB_MCS/Summer26/RobustML(DirectedStudies)/backups"
    os.makedirs(backup_dir, exist_ok=True)
    backup_path = os.path.join(backup_dir, f"Model_Extraction_Attacks_Presentation_backup_{timestamp}.pptx")
    shutil.copy2(output_path, backup_path)
    print(f"Backed up existing presentation to: {backup_path}")
    
    # Load existing file to preserve user's manual modifications
    prs = pptx.Presentation(output_path)
    print("Loaded existing presentation to edit.")
else:
    # Fallback to template if output file does not exist
    prs = pptx.Presentation(template_path)
    print("Output file not found. Initializing from base template.")

# -------------------------------------------------------------
# 2. Deletion function to cleanly remove a slide XML structure
# -------------------------------------------------------------
def clean_delete_slide(prs, index):
    slide_id = prs.slides._sldIdLst[index]
    rId = slide_id.rId
    prs.part.drop_rel(rId)
    prs.slides._sldIdLst.remove(slide_id)

# -------------------------------------------------------------
# 3. Formatting helper for subscripts, superscripts, and math runs
# -------------------------------------------------------------
def populate_formatted_runs(p, text_str, level):
    # Regex to capture:
    # 1. **bold** -> \*\*.*?\*\*
    # 2. *italic* -> \*.*?\*
    # 3. _subscript -> _[a-zA-Z0-9+\-:=]+
    # 4. ^superscript -> \^[a-zA-Z0-9+\-:=()]+
    # 5. Normal text -> anything else
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|_[a-zA-Z0-9+\-:=]+|\^[a-zA-Z0-9+\-:=()]+|[^_*^]+)')
    tokens = pattern.findall(text_str)
    
    for token in tokens:
        run = p.add_run()
        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run.text = token[1:-1]
            run.font.italic = True
        elif token.startswith('_'):
            run.text = token[1:]
            run.font.subscript = True
        elif token.startswith('^'):
            run.text = token[1:]
            run.font.superscript = True
        else:
            run.text = token
            
        run.font.name = "Cambria"
        run.font.size = Pt(18 - 2 * level)
        run.font.color.rgb = RGBColor(0, 0, 0)

# -------------------------------------------------------------
# 4. Insert or Update "Foundational Terminology & Symbols" Slide
# -------------------------------------------------------------
def insert_terminology_slide(prs):
    # Search if terminology slide already exists in presentation
    existing_idx = None
    for idx, slide in enumerate(prs.slides):
        if slide.shapes.title and slide.shapes.title.text == "Foundational Terminology & Symbols":
            existing_idx = idx
            break
            
    if existing_idx is not None:
        clean_delete_slide(prs, existing_idx)
        insert_idx = existing_idx
        print(f"Found existing Terminology slide at index {existing_idx}. Cleanly deleting to update it.")
    else:
        insert_idx = 3 # Insert at index 3 (4th slide position, before current slide 4/5)
        print(f"Inserting new Terminology slide at index {insert_idx}.")
        
    # Add new content slide
    content_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(content_layout)
    
    # Reorder slide to insert_idx
    sldIdLst = prs.slides._sldIdLst
    slide_id_element = sldIdLst[-1]
    sldIdLst.remove(slide_id_element)
    sldIdLst.insert(insert_idx, slide_id_element)
    
    # Title
    title_shape = slide.shapes.title
    title_shape.text = "Foundational Terminology & Symbols"
    for p in title_shape.text_frame.paragraphs:
        for run in p.runs:
            run.font.name = "Cambria"
            run.font.size = Pt(28)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 63, 110) # 003F6E
            
    # Left Column: Terminology (Resize default placeholder)
    body_shape = slide.placeholders[1]
    body_shape.left = int(914400 * 0.8)   # 0.8 inches
    body_shape.width = int(914400 * 4.2)  # 4.2 inches
    body_shape.top = int(914400 * 1.8)    # 1.8 inches
    body_shape.height = int(914400 * 4.8) # 4.8 inches
    tf_left = body_shape.text_frame
    tf_left.word_wrap = True
    
    left_bullets = [
        {"text": "**Core Terminology:**", "level": 0},
        {"text": "• **Target Model** (*f*_[T]): The private victim model hosted on MLaaS or edge devices.", "level": 1},
        {"text": "• **Surrogate Model** (*f*_[S]): The student copycat model trained by the attacker.", "level": 1},
        {"text": "• **Oracle**: The victim API query interface returning predictions/probabilities.", "level": 1},
        {"text": "• **Query** (*x*): An individual input data point sent to the target model.", "level": 1},
        {"text": "• **Query Budget** (*B*): The maximum number of API queries allowed.", "level": 1}
    ]
    
    first = True
    for b in left_bullets:
        if first:
            p = tf_left.paragraphs[0]
            first = False
        else:
            p = tf_left.add_paragraph()
        p.level = b.get("level", 0)
        p.space_after = Pt(6)
        populate_formatted_runs(p, b["text"], p.level)
        
    # Right Column: Symbols & Notation (Add programmatic text box)
    right_left = int(914400 * 5.2)
    right_top = int(914400 * 1.8)
    right_width = int(914400 * 4.0)
    right_height = int(914400 * 4.8)
    right_box = slide.shapes.add_textbox(right_left, right_top, right_width, right_height)
    tf_right = right_box.text_frame
    tf_right.word_wrap = True
    
    right_bullets = [
        {"text": "**Mathematical Symbols:**", "level": 0},
        {"text": "• θ_[T], θ_[S]: Weight parameters of the target and surrogate models.", "level": 1},
        {"text": "• *x* ∈ R^*d*: Input sample feature vector of dimension *d*.", "level": 1},
        {"text": "• *y* ∈ [0, 1]^*c*: Output probability vector for *c* classes.", "level": 1},
        {"text": "• *L*: Loss function minimized during surrogate training.", "level": 1},
        {"text": "• σ(*t*): Logistic sigmoid function: 1 / (1 + *e*^-t).", "level": 1},
        {"text": "• ∇*f*(*x*): Gradient of the function *f* at input *x*.", "level": 1}
    ]
    
    first = True
    for b in right_bullets:
        if first:
            p = tf_right.paragraphs[0]
            first = False
        else:
            p = tf_right.add_paragraph()
        p.level = b.get("level", 0)
        p.space_after = Pt(6)
        populate_formatted_runs(p, b["text"], p.level)

# Run the modification
insert_terminology_slide(prs)

# Save
prs.save(output_path)
print(f"Presentation successfully updated and saved to {output_path}")
